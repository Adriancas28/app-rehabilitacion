"""Genera docs/Capturas_Sistema_SANNA.docx con todas las capturas por rol.

Lee las tablas de docs/analisis/01_Pantallas_y_Capturas.md (ID, pantalla, qué
representa, archivo) y agrega las capturas nuevas de docs/capturas/nuevas_guia_visual
descritas en NUEVAS (abajo). Uso: python docs/generar_capturas_word.py
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

RAIZ = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(RAIZ, "analisis", "01_Pantallas_y_Capturas.md")
CAPS = os.path.join(RAIZ, "capturas")
SALIDA = os.path.join(RAIZ, "Capturas_Sistema_SANNA.docx")

ROLES = [("C", "Pantallas comunes"), ("P", "Rol Paciente"), ("F", "Rol Fisioterapeuta"), ("A", "Rol Administrador")]

# Capturas nuevas (guía visual en vivo): (rol, archivo, título, descripción)
NUEVAS = []
if os.path.exists(os.path.join(RAIZ, "capturas", "nuevas_guia_visual", "descripciones.txt")):
    for linea in open(os.path.join(RAIZ, "capturas", "nuevas_guia_visual", "descripciones.txt"), encoding="utf-8"):
        partes = [p.strip() for p in linea.split("|")]
        if len(partes) == 4:
            NUEVAS.append(tuple(partes))


def filas_md():
    out = []
    for linea in open(MD, encoding="utf-8"):
        if not linea.startswith("| "):
            continue
        celdas = [c.strip() for c in linea.strip().strip("|").split("|")]
        if not re.match(r"^[CPFA]\d+", celdas[0]):
            continue
        archivos = re.findall(r"`([^`]+\.png)`", celdas[-1])
        out.append((celdas[0], celdas[1], celdas[-3].replace("`", ""), archivos))
    return out


def agregar(doc, titulo, descripcion, ruta):
    tabla = doc.add_table(rows=1, cols=2)
    izq, der = tabla.rows[0].cells
    izq.width, der.width = Cm(6.2), Cm(9.8)
    izq.paragraphs[0].add_run().add_picture(ruta, width=Cm(5.6))
    p = der.paragraphs[0]
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(12)
    der.add_paragraph(descripcion).runs[0].font.size = Pt(10)
    doc.add_paragraph()


def main():
    doc = Document()
    doc.add_heading("Capturas de pantalla del sistema SANNA", 0)
    doc.add_paragraph(
        "Todas las pantallas de la aplicación, agrupadas por rol, para elaborar o actualizar la guía de usuario."
    )
    filas = filas_md()
    por_id = {f[0]: f for f in filas}
    carpetas = {"C": "comun", "P": "paciente", "F": "fisio", "A": "admin"}
    total = 0
    for prefijo, nombre in ROLES:
        doc.add_heading(nombre, 1)
        carpeta = carpetas[prefijo]
        ultimo = ("", "")
        for arch in sorted(f for f in os.listdir(os.path.join(CAPS, carpeta)) if f.endswith(".png")):
            m = re.match(r"^([CPFA]\d+)([a-z]?)_(.*)\.png$", arch)
            id_ = m.group(1) if m else arch
            fila = por_id.get(id_)
            if fila:
                ultimo = (fila[1], fila[2])
            titulo = f"{id_}{m.group(2) if m else ''} — " + (m.group(3).replace("_", " ").capitalize() if m else arch)
            desc = ultimo[1] if fila or ultimo[1] else ""
            if not fila and ultimo[1]:
                desc = "Variante de la pantalla " + ultimo[0] + ". " + ultimo[1]
            agregar(doc, titulo, desc, os.path.join(CAPS, carpeta, arch))
            total += 1
        for rol, arch, titulo, desc in NUEVAS:
            if rol == prefijo:
                ruta = os.path.join(CAPS, "nuevas_guia_visual", arch)
                if os.path.exists(ruta):
                    agregar(doc, titulo, desc, ruta)
                    total += 1
    doc.save(SALIDA)
    print("capturas incluidas:", total, "->", SALIDA)


if __name__ == "__main__":
    main()
